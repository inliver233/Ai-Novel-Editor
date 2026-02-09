from __future__ import annotations

import logging
import os
from typing import Dict, List, Tuple

from ...import_export_engine import ExportResult, FormatHandler, ImportResult

logger = logging.getLogger(__name__)


class ExcelHandler(FormatHandler):
    """Excel格式处理器"""

    def can_handle(self, file_path: str) -> bool:
        return file_path.lower().endswith((".xlsx", ".xls"))

    def export_data(self, data: List[Dict], file_path: str, **options) -> ExportResult:
        """导出Excel数据"""
        result = ExportResult(success=False)

        try:
            # 检查是否安装了openpyxl
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Font, PatternFill
            except ImportError:
                result.errors.append("需要安装openpyxl库：pip install openpyxl")
                return result

            self._report_status("准备导出Excel数据...")

            if not data:
                result.errors.append("没有数据可导出")
                return result

            wb = Workbook()
            ws = wb.active
            ws.title = "Codex数据"

            self._report_progress(20, "写入Excel标题...")

            all_fields = set()
            for entry in data:
                all_fields.update(entry.keys())

            priority_fields = ["id", "title", "entry_type", "description", "is_global", "track_references"]
            ordered_fields: List[str] = []

            for field in priority_fields:
                if field in all_fields:
                    ordered_fields.append(field)
                    all_fields.remove(field)

            ordered_fields.extend(sorted(all_fields))

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")

            for col_num, field in enumerate(ordered_fields, 1):
                cell = ws.cell(row=1, column=col_num, value=field)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            self._report_progress(40, "写入Excel数据...")

            for row_num, entry in enumerate(data, 2):
                for col_num, field in enumerate(ordered_fields, 1):
                    value = entry.get(field, "")

                    if isinstance(value, (list, dict)):
                        if field == "aliases" and isinstance(value, list):
                            value = ", ".join(value)
                        else:
                            value = str(value)
                    elif isinstance(value, bool):
                        value = "是" if value else "否"
                    else:
                        value = str(value) if value is not None else ""

                    ws.cell(row=row_num, column=col_num, value=value)

                progress = 40 + (row_num / len(data)) * 50
                self._report_progress(int(progress))

            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            self._report_progress(95, "保存Excel文件...")

            wb.save(file_path)

            self._report_progress(100, "Excel导出完成")

            file_size = os.path.getsize(file_path)
            result.success = True
            result.exported_count = len(data)
            result.file_path = file_path
            result.file_size = file_size

        except Exception as e:
            error_msg = f"Excel导出失败: {e}"
            result.errors.append(error_msg)
            logger.error(error_msg)

        return result

    def import_data(self, file_path: str, **options) -> Tuple[List[Dict], ImportResult]:
        """导入Excel数据"""
        result = ImportResult(success=False)
        data: List[Dict] = []

        try:
            try:
                import openpyxl
            except ImportError:
                result.errors.append("需要安装openpyxl库：pip install openpyxl")
                return data, result

            self._report_status("读取Excel文件...")

            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active

            self._report_progress(30, "解析Excel数据...")

            headers: List[str] = []
            for cell in ws[1]:
                if cell.value:
                    headers.append(str(cell.value))
                else:
                    break

            if not headers:
                raise ValueError("Excel文件中没有找到标题行")

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                if not any(row):
                    continue

                entry: Dict[str, object] = {}
                for col_num, value in enumerate(row):
                    if col_num >= len(headers):
                        break

                    field = headers[col_num]

                    if field in ["is_global", "track_references"] and value:
                        entry[field] = str(value).lower() in ("true", "1", "yes", "是", "true")
                    elif field == "aliases" and value:
                        entry[field] = [alias.strip() for alias in str(value).split(",") if alias.strip()]
                    else:
                        entry[field] = str(value) if value is not None else ""

                data.append(entry)  # type: ignore[arg-type]

                if row_num % 10 == 0:
                    progress = 30 + (row_num / ws.max_row) * 40
                    self._report_progress(int(progress))

            wb.close()

            self._report_progress(80, "验证数据...")

            validation_result = self.validator.validate_codex_data(data)
            result.validation_result = validation_result

            if not validation_result.is_valid and options.get("auto_fix", False):
                self._report_status("自动修复数据...")
                data, fix_log = self.validator.auto_fix_data(data)
                validation_result.fixed_issues.extend(fix_log)
                validation_result = self.validator.validate_codex_data(data)
                result.validation_result = validation_result

            self._report_progress(100, "Excel导入完成")

            result.success = validation_result.is_valid
            result.imported_count = len(data) if result.success else 0
            result.error_count = len(validation_result.errors)

        except Exception as e:
            error_msg = f"Excel导入失败: {e}"
            result.errors.append(error_msg)
            logger.error(error_msg)

        return data, result

