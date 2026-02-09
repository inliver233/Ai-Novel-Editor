from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Tuple

from core.project import DocumentType, ProjectData, ProjectManager

from .markdown import ProgressCallback
from .traversal import collect_novel_documents
from .utils import ensure_novel_root, remove_existing_novel_documents, strip_auto_numbering


def export_project_to_docx(
    project: ProjectData,
    output_path: Path,
    *,
    include_metadata: bool = True,
    title: Optional[str] = None,
    author: Optional[str] = None,
    progress_cb: Optional[ProgressCallback] = None,
) -> int:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:
        raise ImportError("需要安装python-docx库: pip install python-docx") from exc

    documents = collect_novel_documents(project)
    total = len(documents)

    doc = Document()

    if include_metadata:
        resolved_title = title or project.name
        resolved_author = author or project.author

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(resolved_title)
        title_run.font.size = Pt(24)
        title_run.bold = True

        author_para = doc.add_paragraph()
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_para.add_run(f"作者：{resolved_author}")
        author_run.font.size = Pt(14)

        doc.add_page_break()

    for idx, document in enumerate(documents, 1):
        if progress_cb:
            progress_cb(idx, total)

        if document.doc_type == DocumentType.ACT:
            doc.add_heading(f"第{document.order + 1}幕 {document.name}", level=1)
        elif document.doc_type == DocumentType.CHAPTER:
            doc.add_heading(f"第{document.order + 1}章 {document.name}", level=2)
        elif document.doc_type == DocumentType.SCENE:
            doc.add_heading(f"场景{document.order + 1}：{document.name}", level=3)

        if document.content:
            paragraphs = document.content.split("\n")
            for para_text in paragraphs:
                if para_text.strip():
                    para = doc.add_paragraph(para_text)
                    para.paragraph_format.first_line_indent = Pt(24)

        if document.doc_type in {DocumentType.ACT, DocumentType.CHAPTER} and idx < len(documents):
            doc.add_page_break()

    doc.save(str(output_path))
    return total


def import_novel_from_docx(
    manager: ProjectManager,
    input_path: Path,
    *,
    replace_existing: bool = False,
    progress_cb: Optional[ProgressCallback] = None,
) -> int:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("需要安装python-docx库: pip install python-docx") from exc

    doc = Document(str(input_path))

    sections: List[Tuple[int, str, List[str]]] = []
    current_section: Optional[Tuple[int, str, List[str]]] = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style.name.startswith("Heading"):
            if current_section:
                sections.append(current_section)

            level = int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1
            current_section = (level, text, [])
        else:
            if current_section:
                current_section[2].append(text)
            else:
                current_section = (1, "导入内容", [text])

    if current_section:
        sections.append(current_section)

    novel_root_id = ensure_novel_root(manager)
    if replace_existing:
        remove_existing_novel_documents(manager)

    total = len(sections)
    imported_count = 0
    parent_map: dict[int, str] = {}

    for idx, (level, raw_title, paragraphs) in enumerate(sections, 1):
        if progress_cb:
            progress_cb(idx, total)

        content = "\n\n".join(paragraphs).strip()

        if level <= 1:
            doc_type = DocumentType.ACT
            parent_id: Optional[str] = novel_root_id
            parent_map = {}
        elif level == 2:
            doc_type = DocumentType.CHAPTER
            parent_id = parent_map.get(1) or novel_root_id
        else:
            doc_type = DocumentType.SCENE
            parent_id = parent_map.get(2) or parent_map.get(1) or novel_root_id

        title = strip_auto_numbering(raw_title, doc_type=doc_type)
        created = manager.add_document(title, doc_type, parent_id, save=False)
        if not created:
            continue
        if content:
            manager.update_document(created.id, content=content, save=False)

        if doc_type == DocumentType.ACT:
            parent_map[1] = created.id
        elif doc_type == DocumentType.CHAPTER:
            parent_map[2] = created.id

        imported_count += 1

    manager.save_project()
    return imported_count
